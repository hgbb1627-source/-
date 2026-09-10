# -*- coding: utf-8 -*-
"""
수시 실시간 경쟁률 자동 수집 및 동기화 시스템 (Auto Fetch Competition Tool)
- 11개 대상 전형(지원 6개, 모니터링 5개)의 실시간 경쟁률 웹페이지(유웨이/진학사)를 자동 파싱
- 새로운 발표 기준시각 또는 인원 변동 시 엑셀 파일(로컬 및 Google Drive)에 자동 행 추가
- 웹앱용 정형 데이터(data/records.json) 자동 추출 및 동기화
- README.md 종합비교 대시보드 실시간 테이블 자동 갱신
- GitHub 자동 커밋 및 푸시 지원

사용법:
  1) 단순 조회만 수행 (Dry-run):
     python auto_fetch_competition.py --check
  2) 엑셀, JSON 및 README에 저장:
     python auto_fetch_competition.py --save
  3) 엑셀, JSON, README 저장 및 GitHub 커밋·푸시:
     python auto_fetch_competition.py --save --push
"""

import os
import sys
import io
import re
import json
import argparse
import datetime
import shutil
import urllib.request
import http.cookiejar
from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import Font, Alignment

# UTF-8 콘솔 출력 보장
if sys.platform == "win32":
    try:
        sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
        sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8", errors="replace")
    except Exception:
        pass

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
LOCAL_EXCEL_PATH = os.path.join(BASE_DIR, "data", "2027 수시 경쟁률 예측 프로그램.xlsx")
GDRIVE_EXCEL_PATH = r"G:\내 드라이브\02. My Documents\2027학년도 대입\2027 수시 경쟁률 예측 프로그램.xlsx"
README_PATH = os.path.join(BASE_DIR, "README.md")
JSON_PATH = os.path.join(BASE_DIR, "data", "records.json")

def get_excel_path():
    if os.path.exists(GDRIVE_EXCEL_PATH):
        return GDRIVE_EXCEL_PATH
    elif os.path.exists(LOCAL_EXCEL_PATH):
        return LOCAL_EXCEL_PATH
    return GDRIVE_EXCEL_PATH

# 11개 대상 전형 설정 (의약학 계열: 핵심 6 + 모니터링 5)
# ※ 전북대/원광대/가천대는 하나의 실시간 경쟁률 URL에 여러 전형이 함께 표시되므로,
#   alias/sheet를 대학명+전형 구분자로 "유일하게" 만들어야 함(동일 alias 사용 시 딕셔너리 키 충돌로 데이터 덮어써짐).
UNIV_CONFIGS = [
    {
        "category": "지원 1",
        "alias": "전북대-지역의사전주권",
        "sheet": "전북대_지역의사전주권",
        "univ_name": "전북대학교",
        "admission": "학생부교과 지역의사선발전형-전주권",
        "major": "의예과",
        "quota": 9,
        "prev_year": "신설",
        "url": "https://ratio.uwayapply.com/Sl5KOldCL0pmJSY6Jko3ZlRm",
        "enc": "euc-kr",
        "type": "uway",
        "target_admission_match": "지역의사선발전형-전주권",
        "target_major_match": "의예과"
    },
    {
        "category": "지원 2",
        "alias": "전북대-지역인재2전북권",
        "sheet": "전북대_지역인재2전북권",
        "univ_name": "전북대학교",
        "admission": "학생부교과 지역인재2전형-전북권",
        "major": "의예과",
        "quota": 45,
        "prev_year": "10.0 : 1",
        "url": "https://ratio.uwayapply.com/Sl5KOldCL0pmJSY6Jko3ZlRm",
        "enc": "euc-kr",
        "type": "uway",
        "target_admission_match": "지역인재2전형-전북권",
        "target_major_match": "의예과"
    },
    {
        "category": "지원 3",
        "alias": "가천대-의예과",
        "sheet": "가천대_의예과",
        "univ_name": "가천대학교",
        "admission": "가천의약학전형",
        "major": "의예과",
        "quota": 16,
        "prev_year": "40.0 : 1",
        "url": "https://addon.jinhakapply.com/RatioV1/RatioH/Ratio10190711.html",
        "enc": "utf-8",
        "type": "jinhak",
        "target_admission_match": "가천의약학전형",
        "target_major_match": "의예과"
    },
    {
        "category": "지원 4",
        "alias": "원광대-의약학치의예과",
        "sheet": "원광대_의약학치의예과",
        "univ_name": "원광대학교",
        "admission": "지역인재종합전형(전북)",
        "major": "의약학 치의예과 (자연)",
        "quota": 18,
        "prev_year": "6.1 : 1",
        "url": "https://ratio.uwayapply.com/Sl5Kclc4TjlXYU5KZiUmOiZKN2ZUZg==",
        "enc": "euc-kr",
        "type": "uway",
        "target_admission_match": "지역인재종합전형",
        "target_major_match": "치의예과"
    },
    {
        "category": "지원 5",
        "alias": "성균관대-약학과",
        "sheet": "성균관대_약학과",
        "univ_name": "성균관대학교",
        "admission": "학생부종합(융합인재)",
        "major": "약학과",
        "quota": 10,
        "prev_year": "신설",
        "url": "https://addon.jinhakapply.com/RatioV1/RatioH/Ratio10920591.html",
        "enc": "utf-8",
        "type": "jinhak",
        "target_admission_match": "융합인재",
        "target_major_match": "약학과"
    },
    {
        "category": "지원 6",
        "alias": "한양대ERICA-약학과",
        "sheet": "한양대ERICA_약학과",
        "univ_name": "한양대학교(ERICA)",
        "admission": "학생부종합(서류형)",
        "major": "약학과",
        "quota": 15,
        "prev_year": "35.7 : 1",
        "url": "https://addon.jinhakapply.com/RatioV1/RatioH/Ratio11650731.html",
        "enc": "utf-8",
        "type": "jinhak",
        "target_admission_match": "서류형",
        "target_major_match": "약학과"
    },
    {
        "category": "기타",
        "alias": "전북대-지역인재2유형",
        "sheet": "전북대_지역인재2유형",
        "univ_name": "전북대학교",
        "admission": "학생부종합 지역인재2유형전형-전북권",
        "major": "의예과",
        "quota": 4,
        "prev_year": "신설",
        "url": "https://ratio.uwayapply.com/Sl5KOldCL0pmJSY6Jko3ZlRm",
        "enc": "euc-kr",
        "type": "uway",
        "target_admission_match": "지역인재2유형전형-전북권",
        "target_major_match": "의예과"
    },
    {
        "category": "기타",
        "alias": "전북대-지역의사광역권",
        "sheet": "전북대_지역의사광역권",
        "univ_name": "전북대학교",
        "admission": "학생부교과 지역의사선발전형-광역권",
        "major": "의예과",
        "quota": 6,
        "prev_year": "신설",
        "url": "https://ratio.uwayapply.com/Sl5KOldCL0pmJSY6Jko3ZlRm",
        "enc": "euc-kr",
        "type": "uway",
        "target_admission_match": "지역의사선발전형-광역권",
        "target_major_match": "의예과"
    },
    {
        "category": "기타",
        "alias": "원광대-지역의사전주권",
        "sheet": "원광대_지역의사전주권",
        "univ_name": "원광대학교",
        "admission": "지역의사선발전형 진료권-전주권",
        "major": "의예과",
        "quota": 7,
        "prev_year": "신설",
        "url": "https://ratio.uwayapply.com/Sl5Kclc4TjlXYU5KZiUmOiZKN2ZUZg==",
        "enc": "euc-kr",
        "type": "uway",
        "target_admission_match": "지역의사선발전형",
        "target_major_match": "의예과"
    },
    {
        "category": "기타",
        "alias": "원광대-지역의사광역권",
        "sheet": "원광대_지역의사광역권",
        "univ_name": "원광대학교",
        "admission": "지역의사선발전형 진료권-광역권",
        "major": "의예과",
        "quota": 4,
        "prev_year": "신설",
        "url": "https://ratio.uwayapply.com/Sl5Kclc4TjlXYU5KZiUmOiZKN2ZUZg==",
        "enc": "euc-kr",
        "type": "uway",
        "target_admission_match": "지역의사선발전형",
        "target_major_match": "의예과"
    },
    {
        "category": "기타",
        "alias": "가천대-약학과",
        "sheet": "가천대_약학과",
        "univ_name": "가천대학교",
        "admission": "가천의약학전형",
        "major": "약학과",
        "quota": 10,
        "prev_year": "50.0 : 1",
        "url": "https://addon.jinhakapply.com/RatioV1/RatioH/Ratio10190711.html",
        "enc": "utf-8",
        "type": "jinhak",
        "target_admission_match": "가천의약학전형",
        "target_major_match": "약학과"
    }
]

def parse_time_from_soup(soup):
    """HTML 내용 전체에서 기준 시각(날짜 및 시간)을 정규식으로 안전하게 추출"""
    pattern = r'(\d{4})[-\.년]\s*(\d{1,2})[-\.월]\s*(\d{1,2})[일\.]?\s*(오전|오후)?\s*(\d{1,2})[:시]\s*(\d{1,2})'
    for s in soup.stripped_strings:
        m = re.search(pattern, s)
        if m:
            y, mo, d, ampm, hh, mm = m.groups()
            hh, mm = int(hh), int(mm)
            if ampm == "오후" and hh < 12:
                hh += 12
            elif ampm == "오전" and hh == 12:
                hh = 0
            date_obj = datetime.date(int(y), int(mo), int(d))
            time_obj = datetime.time(hh, mm)
            return date_obj, time_obj
    return None, None

_BROWSER_UA = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"

def _fetch_html_with_session(url, timeout=10):
    """쿠키 세션을 유지한 채로 대상 사이트 루트를 먼저 방문(세션 쿠키 확보) 후
    실제 페이지를 요청 — 진학사 등 세션/쿠키 기반 봇 차단 회피 시도"""
    domain_root = "https://" + url.split("/")[2] + "/"
    common_headers = {
        "User-Agent": _BROWSER_UA,
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
        "Accept-Language": "ko-KR,ko;q=0.9,en-US;q=0.8,en;q=0.7",
        "Connection": "keep-alive",
        "Upgrade-Insecure-Requests": "1"
    }
    cookie_jar = http.cookiejar.CookieJar()
    opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(cookie_jar))

    try:
        warmup_req = urllib.request.Request(domain_root, headers=common_headers)
        opener.open(warmup_req, timeout=timeout)
    except Exception:
        pass  # 워밍업 실패해도 본 요청은 시도

    req = urllib.request.Request(url, headers={**common_headers, "Referer": domain_root})
    with opener.open(req, timeout=timeout) as resp:
        return resp.read()

def fetch_single_university(cfg):
    """단일 대학 웹페이지를 스크래핑하여 최신 수치 딕셔너리 반환"""
    try:
        raw = _fetch_html_with_session(cfg["url"])
        try:
            html = raw.decode(cfg["enc"])
        except Exception:
            html = raw.decode("euc-kr", errors="replace")

        soup = BeautifulSoup(html, "html.parser")
        
        date_obj, time_obj = parse_time_from_soup(soup)
        
        target_adm = cfg["target_admission_match"]
        target_maj = cfg["target_major_match"]
        
        matched_row = None
        for table in soup.find_all("table"):
            headers = []
            curr = table
            while curr and len(headers) < 4:
                curr = curr.find_previous(["h1", "h2", "h3", "h4", "caption", "div", "span", "p"])
                if curr:
                    t = curr.get_text().strip()
                    if any(k in t for k in ["전형", "모집", "현황", "경쟁률"]) and len(t) < 60:
                        headers.append(t)
            header_text = " ".join(headers)

            if target_adm and target_adm not in header_text:
                table_first_col = " ".join(
                    [tr.find(["td", "th"]).get_text() for tr in table.find_all("tr") if tr.find(["td", "th"])]
                )
                if target_adm not in table_first_col:
                    continue

            for tr in table.find_all("tr"):
                cells = [td.get_text().strip() for td in tr.find_all(["td", "th"])]
                row_str = " ".join(cells)
                if target_maj in row_str:
                    nums = [c for c in cells if c.isdigit()]
                    if len(nums) >= 2:
                        q_str = str(cfg["quota"])
                        if q_str in cells:
                            q_idx = cells.index(q_str)
                        else:
                            q_idx = cells.index(nums[-2])
                        
                        if q_idx + 1 < len(cells) and cells[q_idx + 1].isdigit():
                            app = int(cells[q_idx + 1])
                            rate_str = cells[q_idx + 2] if q_idx + 2 < len(cells) else f"{app / cfg['quota']:.1f} : 1"
                            matched_row = {
                                "applicants": app,
                                "rate_str": rate_str
                            }
                            break
            if matched_row:
                break

        # 1차 매칭 실패 시: 전형명(target_adm)으로 표를 좁히지 않고, 모집단위명 + 정원 숫자
        # 일치만으로 전체 표를 다시 훑는 완화된 2차 매칭 (한 URL에 여러 전형이 섞여 있어
        # 전형명 문구가 실제 페이지와 미세하게 달라 표 단계에서 걸러지는 경우를 구제)
        if not matched_row:
            q_str = str(cfg["quota"])
            for table in soup.find_all("table"):
                for tr in table.find_all("tr"):
                    cells = [td.get_text().strip() for td in tr.find_all(["td", "th"])]
                    row_str = " ".join(cells)
                    if target_maj in row_str and q_str in cells:
                        q_idx = cells.index(q_str)
                        if q_idx + 1 < len(cells) and cells[q_idx + 1].isdigit():
                            app = int(cells[q_idx + 1])
                            rate_str = cells[q_idx + 2] if q_idx + 2 < len(cells) else f"{app / cfg['quota']:.1f} : 1"
                            matched_row = {"applicants": app, "rate_str": rate_str}
                            break
                if matched_row:
                    break
            if matched_row:
                print(f"[알림] [{cfg['alias']}] 완화된 2차 매칭으로 찾음 (target_admission_match 확인 권장)")

        if not matched_row:
            print(f"[경고] [{cfg['alias']}] 타겟 학과/전형 행 매칭 실패")
            return None

        return {
            "cfg": cfg,
            "date": date_obj or datetime.date.today(),
            "time": time_obj or datetime.time(datetime.datetime.now().hour, 0),
            "applicants": matched_row["applicants"],
            "rate_str": matched_row["rate_str"]
        }

    except Exception as e:
        print(f"[오류] [{cfg['alias']}] 웹페이지 요청 또는 파싱 실패: {e}")
        return None

def get_current_excel_last_records():
    """엑셀 파일에서 각 대학 시트의 가장 최근 기록 행 정보 읽기"""
    excel_path = get_excel_path()
    if not os.path.exists(excel_path):
        return {}

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    last_records = {}
    for cfg in UNIV_CONFIGS:
        sname = cfg["sheet"]
        if sname not in wb.sheetnames:
            continue
        ws = wb[sname]
        last_r = None
        for r in range(21, 81):
            if ws.cell(r, 4).value is not None:
                last_r = r
        if last_r:
            d_val = ws.cell(last_r, 2).value
            t_val = ws.cell(last_r, 3).value
            cnt_val = ws.cell(last_r, 4).value
            if isinstance(d_val, datetime.datetime):
                d_val = d_val.date()
            if isinstance(t_val, datetime.datetime):
                t_val = t_val.time()
            last_records[cfg["alias"]] = {
                "row": last_r,
                "date": d_val,
                "time": t_val,
                "applicants": cnt_val
            }
        else:
            last_records[cfg["alias"]] = None
    return last_records

def get_current_json_last_records():
    """data/records.json에서 각 대학의 가장 최근 기록 읽기 (엑셀이 없는 환경용 폴백)"""
    if not os.path.exists(JSON_PATH):
        return {}
    try:
        with open(JSON_PATH, "r", encoding="utf-8") as f:
            existing = json.load(f)
    except Exception:
        return {}
    last_records = {}
    for u in existing.get("universities", []):
        history = u.get("history", [])
        if not history:
            continue
        last = history[-1]
        try:
            d_obj = datetime.datetime.strptime(last["date"], "%Y-%m-%d").date()
            t_obj = datetime.datetime.strptime(last["time"], "%H:%M").time()
        except Exception:
            continue
        last_records[u["alias"]] = {
            "date": d_obj,
            "time": t_obj,
            "applicants": last.get("applicants")
        }
    return last_records

def fetch_all():
    """대상 전형 전체 데이터 수집 (웹 스크래핑 실패 시 엑셀 또는 직전 records.json 데이터로 Fallback 유지)"""
    print("=" * 80)
    print(f" [수집 시작] {len(UNIV_CONFIGS)}개 대상 전형 실시간 경쟁률 웹 스크래핑 진행 중...")
    print("=" * 80)
    excel_records = get_current_excel_last_records()
    json_records = get_current_json_last_records()
    results = []
    for cfg in UNIV_CONFIGS:
        res = fetch_single_university(cfg)
        if res:
            results.append(res)
            d_s = res["date"].strftime("%Y-%m-%d")
            t_s = res["time"].strftime("%H:%M")
            quota = cfg["quota"]
            app = res["applicants"]
            rate = round(app / quota, 1)
            print(f"  ✓ [{cfg['alias']:<6}] {d_s} {t_s} | 모집: {quota:>2}명 | 지원자: {app:>4}명 | 경쟁률: {rate:.1f} : 1")
        else:
            # Fallback: 웹 스크래핑 실패 시(해외 IP 차단 등) 엑셀 또는 직전 records.json 기록 유지
            last = excel_records.get(cfg["alias"]) or json_records.get(cfg["alias"])
            if last and last["applicants"] is not None:
                d_obj = last["date"] if isinstance(last["date"], datetime.date) else datetime.date.today()
                t_obj = last["time"] if isinstance(last["time"], datetime.time) else datetime.time(0, 0)
                app_cnt = int(last["applicants"])
                fallback_res = {
                    "cfg": cfg,
                    "date": d_obj,
                    "time": t_obj,
                    "applicants": app_cnt,
                    "rate_str": f"{app_cnt / cfg['quota']:.1f} : 1",
                    "is_fallback": True
                }
                results.append(fallback_res)
                d_s = d_obj.strftime("%Y-%m-%d")
                t_s = t_obj.strftime("%H:%M")
                quota = cfg["quota"]
                rate = round(app_cnt / quota, 1)
                print(f"  [유지] [{cfg['alias']:<6}] 웹 스크래핑 실패로 직전 기록 유지 ({d_s} {t_s} | 지원자 {app_cnt:>4}명 | 경쟁률 {rate:.1f} : 1)")
            else:
                print(f"  [경고] [{cfg['alias']:<6}] 엑셀/JSON에도 기존 기록이 없어 수집 제외됨")
    return results

def sync_excel_and_record(fetched_list):
    """변경된 대학만 엑셀에 새 행 추가 기록"""
    excel_path = get_excel_path()
    last_records = get_current_excel_last_records()

    wb = openpyxl.load_workbook(excel_path)
    updated_count = 0
    update_summary = []

    for item in fetched_list:
        if item.get("is_fallback"):
            continue
        cfg = item["cfg"]
        alias = cfg["alias"]
        sname = cfg["sheet"]
        cur_d = item["date"]
        cur_t = item["time"]
        cur_app = item["applicants"]

        last = last_records.get(alias)
        need_update = False

        if last is None:
            need_update = True
            prev_app = 0
        else:
            prev_d = last["date"]
            prev_t = last["time"]
            prev_app = last["applicants"] or 0

            time_changed = (prev_d != cur_d) or (prev_t != cur_t)
            app_increased = cur_app > prev_app

            if time_changed or app_increased:
                need_update = True

        if not need_update:
            continue

        ws = wb[sname]
        target_row = None
        for r in range(21, 81):
            if ws.cell(r, 2).value is None and ws.cell(r, 3).value is None and ws.cell(r, 4).value is None:
                target_row = r
                break

        if target_row is None:
            print(f"  [오류] '{sname}' 시트의 기록 영역(21~80행)이 가득 찼습니다.")
            continue

        prev_row = target_row - 1
        ws.cell(target_row, 2).value = cur_d
        ws.cell(target_row, 3).value = cur_t
        ws.cell(target_row, 4).value = int(cur_app)

        ws.cell(target_row, 2).number_format = ws.cell(prev_row, 2).number_format or "yyyy-mm-dd"
        ws.cell(target_row, 3).number_format = ws.cell(prev_row, 3).number_format or "h:mm"
        ws.cell(target_row, 4).number_format = ws.cell(prev_row, 4).number_format or "#,##0"

        for c in range(2, 5):
            curr_cell = ws.cell(target_row, c)
            ref_cell = ws.cell(prev_row, c)
            if ref_cell.has_style:
                curr_cell.font = Font(
                    name=ref_cell.font.name,
                    size=ref_cell.font.size,
                    bold=ref_cell.font.bold,
                    italic=ref_cell.font.italic,
                    color=ref_cell.font.color
                )
                curr_cell.alignment = Alignment(
                    horizontal=ref_cell.alignment.horizontal,
                    vertical=ref_cell.alignment.vertical
                )

        diff = cur_app - prev_app
        rate = round(cur_app / cfg["quota"], 1)
        prev_year_rate = ws["C15"].value or cfg["prev_year"]

        update_summary.append({
            "alias": alias,
            "univ_name": cfg["univ_name"],
            "time_str": cur_t.strftime("%H:%M"),
            "quota": cfg["quota"],
            "applicants": cur_app,
            "diff": diff,
            "rate": rate,
            "prev_year_rate": prev_year_rate
        })
        updated_count += 1

    if updated_count > 0:
        wb.save(excel_path)
        print(f"\n[엑셀 갱신 완료] 총 {updated_count}개 대학의 최신 데이터가 반영되었습니다: {excel_path}")

        other_path = LOCAL_EXCEL_PATH if excel_path == GDRIVE_EXCEL_PATH else GDRIVE_EXCEL_PATH
        if os.path.exists(os.path.dirname(other_path)):
            try:
                shutil.copy2(excel_path, other_path)
                target_name = "프로젝트 로컬" if other_path == LOCAL_EXCEL_PATH else "Google Drive"
                print(f"  [동기화 완료] {target_name} 파일에도 동기화 복사되었습니다: {other_path}")
            except Exception as e:
                print(f"  [알림] 동기화 생략: {e}")
    else:
        print("\n[엑셀 유지] 모든 대학이 이미 최신 상태이거나 변동 사항이 없습니다.")

    return update_summary

def _load_existing_json_history():
    """기존 data/records.json에서 alias별 history를 읽어옴 (엑셀이 없는 클라우드 환경용 폴백)"""
    if not os.path.exists(JSON_PATH):
        return {}
    try:
        with open(JSON_PATH, "r", encoding="utf-8") as f:
            existing = json.load(f)
        return {u["alias"]: u.get("history", []) for u in existing.get("universities", [])}
    except Exception:
        return {}

def export_records_json(fetched_list):
    """웹앱에서 실시간으로 읽을 수 있는 경량 JSON 데이터셋(data/records.json) 추출 및 저장"""
    excel_path = get_excel_path()
    wb = None
    if os.path.exists(excel_path):
        try:
            wb = openpyxl.load_workbook(excel_path, data_only=True)
        except Exception as e:
            print(f"[알림] 엑셀 읽기 실패, JSON 히스토리 폴백 사용: {e}")

    # 엑셀이 없는 환경(GitHub Actions 등)에서는 직전 커밋된 records.json의
    # history를 이어받아 누적한다 (있으면).
    json_history_fallback = _load_existing_json_history() if wb is None else {}

    latest_time_str = "00:00"
    univ_data_list = []

    for item in fetched_list:
        cfg = item["cfg"]
        alias = cfg["alias"]
        sname = cfg["sheet"]
        cur_t_str = item["time"].strftime("%H:%M")
        cur_d_str = item["date"].strftime("%Y-%m-%d")
        if cur_t_str > latest_time_str:
            latest_time_str = cur_t_str

        history = []
        if wb is not None and sname in wb.sheetnames:
            ws = wb[sname]
            for r in range(21, 81):
                d_val = ws.cell(r, 2).value
                t_val = ws.cell(r, 3).value
                app_val = ws.cell(r, 4).value
                if app_val is not None:
                    d_str = d_val.strftime("%Y-%m-%d") if isinstance(d_val, (datetime.date, datetime.datetime)) else str(d_val)
                    t_str = t_val.strftime("%H:%M") if isinstance(t_val, (datetime.time, datetime.datetime)) else str(t_val)
                    q = cfg["quota"]
                    rate = round(int(app_val) / q, 2)
                    history.append({
                        "date": d_str,
                        "time": t_str,
                        "datetime_label": f"{d_str[-5:]} {t_str}",
                        "applicants": int(app_val),
                        "rate": rate
                    })
        else:
            # JSON 히스토리 폴백: 직전 기록을 이어받고, 새 시점/증가분만 추가
            history = list(json_history_fallback.get(alias, []))
            cur_app_val = item["applicants"]
            last_h = history[-1] if history else None
            is_new_point = (not last_h) or (last_h.get("time") != cur_t_str or last_h.get("date") != cur_d_str)
            if is_new_point and (not last_h or cur_app_val >= last_h.get("applicants", 0)):
                cur_rate_val = round(cur_app_val / cfg["quota"], 1)
                history.append({
                    "date": cur_d_str,
                    "time": cur_t_str,
                    "datetime_label": f"{cur_d_str[-5:]} {cur_t_str}",
                    "applicants": cur_app_val,
                    "rate": cur_rate_val
                })

        diff = 0
        if len(history) >= 2:
            diff = history[-1]["applicants"] - history[-2]["applicants"]

        cur_app = item["applicants"]
        cur_rate = round(cur_app / cfg["quota"], 1)

        # 전년도 대비 달성률 (%) — 전년도 경쟁률 데이터가 없으면(예: "데이터 없음") 진행률을 계산하지 않고 None 처리
        progress_pct = None
        if ":" in cfg["prev_year"]:
            prev_yr_num = float(cfg["prev_year"].split(":")[0].strip())
            if prev_yr_num > 0:
                progress_pct = round((cur_rate / prev_yr_num) * 100, 1)

        univ_data_list.append({
            "category": cfg["category"],
            "alias": alias,
            "univ_name": cfg["univ_name"],
            "admission": cfg["admission"],
            "major": cfg["major"],
            "quota": cfg["quota"],
            "applicants": cur_app,
            "diff": diff,
            "rate": cur_rate,
            "rate_str": f"{cur_rate:.1f} : 1",
            "prev_year": cfg["prev_year"],
            "progress_pct": progress_pct,
            "latest_time": cur_t_str,
            "url": cfg["url"],
            "history": history
        })

    payload = {
        "updated_at": f"{datetime.date.today().strftime('%Y-%m-%d')} {latest_time_str}",
        "total_target_colleges": len(univ_data_list),
        "universities": univ_data_list
    }

    os.makedirs(os.path.dirname(JSON_PATH), exist_ok=True)
    with open(JSON_PATH, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)

    print(f"[JSON 갱신 완료] 웹앱 데이터 파일 생성 완료: {JSON_PATH}")

def update_readme(fetched_list):
    """README.md의 종합비교 대시보드 표를 최신 수치로 동기화"""
    if not os.path.exists(README_PATH):
        return

    with open(README_PATH, "r", encoding="utf-8") as f:
        content = f.read()

    latest_time_str = "00:00"
    for item in fetched_list:
        t_str = item["time"].strftime("%H:%M")
        if t_str > latest_time_str:
            latest_time_str = t_str

    now_date_str = datetime.date.today().strftime("%Y-%m-%d")

    rows_text = []
    for item in fetched_list:
        cfg = item["cfg"]
        cat = cfg["category"]
        t_str = item["time"].strftime("%H:%M")
        u_name = cfg["univ_name"]
        adm = cfg["admission"]
        maj = cfg["major"]
        quota = cfg["quota"]
        app = item["applicants"]
        rate = round(app / quota, 1)

        prev_yr_str = cfg["prev_year"]

        if "지원" in cat:
            row_line = f"| **{cat}** | **{t_str}** | **{u_name}** | {adm} | {maj} | **{quota}** | **{app}** | **{rate:.1f} : 1** | {prev_yr_str} |"
        else:
            row_line = f"| {cat} | **{t_str}** | **{u_name}** | {adm} | {maj} | **{quota}** | **{app}** | **{rate:.1f} : 1** | {prev_yr_str} |"
        rows_text.append(row_line)

    new_table = f"### 📊 종합비교 대시보드 실시간 현황 ({now_date_str} {latest_time_str} 기준)\n\n"
    new_table += "| 구분 | 발표시간 | 대학명 | 전형명 | 모집단위 | 모집정원 | 최신 지원자 | 현재 경쟁률 | 전년도 경쟁률 (2026) |\n"
    new_table += "| :---: | :---: | :--- | :--- | :--- | :---: | :---: | :---: | :---: |\n"
    new_table += "\n".join(rows_text) + "\n"

    pattern = r"### 📊 종합비교 대시보드 실시간 현황.*?(?=\n---\n|\n##|\Z)"
    updated_content = re.sub(pattern, new_table, content, flags=re.DOTALL)

    with open(README_PATH, "w", encoding="utf-8") as f:
        f.write(updated_content)

    print(f"[README 갱신 완료] 대시보드 표가 최신 수치({now_date_str} {latest_time_str})로 동기화되었습니다.")

def git_commit_and_push(latest_time_str):
    """Git 변경사항 커밋 및 푸시"""
    import subprocess
    msg = f"Update admission records automatically ({datetime.date.today().strftime('%Y-%m-%d')} {latest_time_str})"
    print(f"\n[Git 동기화] 커밋 및 푸시 진행 중: '{msg}'")
    try:
        subprocess.run(["git", "add", "."], cwd=BASE_DIR, check=True)
        status = subprocess.run(["git", "status", "--porcelain"], cwd=BASE_DIR, capture_output=True, text=True)
        if not status.stdout.strip():
            print("  [Git 알림] 커밋할 변경사항이 없습니다.")
            return
        subprocess.run(["git", "commit", "-m", msg], cwd=BASE_DIR, check=True)
        branch = subprocess.run(
            ["git", "rev-parse", "--abbrev-ref", "HEAD"], cwd=BASE_DIR,
            capture_output=True, text=True, check=True
        ).stdout.strip()
        subprocess.run(["git", "push", "origin", branch], cwd=BASE_DIR, check=True)
        print(f"  ✓ GitHub 푸시 완료! ({branch} 브랜치)")
    except Exception as e:
        print(f"  [Git 오류] Git 작업 실패: {e}")

def git_pull_and_sync():
    """GitHub 최신 데이터를 pull하고 로컬 및 구글 드라이브 엑셀 파일 동기화"""
    import subprocess
    print("\n[Git 동기화] GitHub 최신 데이터 내려받기(Pull) 진행 중...")
    try:
        branch = subprocess.run(
            ["git", "rev-parse", "--abbrev-ref", "HEAD"], cwd=BASE_DIR,
            capture_output=True, text=True, check=True
        ).stdout.strip()
        subprocess.run(["git", "pull", "origin", branch], cwd=BASE_DIR, check=True)
        print(f"  ✓ Git Pull 성공! ({branch} 브랜치)")
        if os.path.exists(LOCAL_EXCEL_PATH) and os.path.exists(os.path.dirname(GDRIVE_EXCEL_PATH)):
            shutil.copy2(LOCAL_EXCEL_PATH, GDRIVE_EXCEL_PATH)
            print(f"  ✓ [동기화 완료] Google Drive 엑셀 파일이 최신 버전으로 갱신되었습니다:\n    -> {GDRIVE_EXCEL_PATH}")
        else:
            print("  [알림] Google Drive 경로를 찾을 수 없어 로컬 파일만 갱신되었습니다.")
    except Exception as e:
        print(f"  [오류] Git Pull 또는 동기화 실패: {e}")

def main():
    parser = argparse.ArgumentParser(description="수시 실시간 경쟁률 자동 수집 및 동기화 도구")
    parser.add_argument("--check", action="store_true", help="수집 결과만 화면에 표시 (저장하지 않음)")
    parser.add_argument("--save", action="store_true", help="엑셀, JSON 및 README에 저장")
    parser.add_argument("--push", action="store_true", help="저장 후 GitHub에 커밋 및 푸시")
    parser.add_argument("--pull", action="store_true", help="GitHub 최신 데이터 Pull 및 구글 드라이브 동기화")
    parser.add_argument("--daemon", type=int, nargs="?", const=5, help="N분 간격으로 실시간 수집·저장·푸시 상시 실행 (기본 5분)")

    args = parser.parse_args()

    if args.pull:
        git_pull_and_sync()
        return

    if args.daemon:
        interval_min = args.daemon
        print(f"\n[데몬 모드 가동] {interval_min}분 간격으로 {len(UNIV_CONFIGS)}개 전형 자동 수집 및 동기화를 시작합니다. (종료: Ctrl+C)")
        import time
        while True:
            try:
                now_str = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
                print(f"\n[{now_str}] 정기 수집 가동...")
                fetched = fetch_all()
                if fetched:
                    try:
                        sync_excel_and_record(fetched)
                    except Exception as e:
                        print(f"[알림] 엑셀 동기화 생략(로컬/구글드라이브 파일 없음 등): {e}")
                    update_readme(fetched)
                    export_records_json(fetched)
                    latest_t = max([item["time"].strftime("%H:%M") for item in fetched])
                    git_commit_and_push(latest_t)
            except Exception as e:
                print(f"  [데몬 오류]: {e}")
            print(f"\n다음 수집까지 {interval_min}분 대기합니다...")
            time.sleep(interval_min * 60)

    if not args.check and not args.save:
        print("옵션을 지정해 주세요: --check, --save, --push, --pull, --daemon [분]")
        print("예시: python auto_fetch_competition.py --save --push")
        sys.exit(0)

    fetched = fetch_all()
    if not fetched:
        print("[오류] 수집된 데이터가 없습니다.")
        sys.exit(1)

    if args.check:
        print("\n[조회 완료] --check 모드이므로 엑셀 및 README를 수정하지 않았습니다.")
        return

    if args.save:
        # 엑셀 동기화는 로컬/구글 드라이브 파일이 있을 때만 의미가 있는 보조 기능입니다.
        # GitHub Actions 등 클라우드 환경에는 그 파일이 없으므로 실패해도 무시하고
        # 웹 대시보드가 실제로 읽는 README/JSON 갱신과 git push는 계속 진행합니다.
        try:
            sync_excel_and_record(fetched)
        except Exception as e:
            print(f"[알림] 엑셀 동기화 생략(로컬/구글드라이브 파일 없음 등): {e}")
        update_readme(fetched)
        export_records_json(fetched)

    if args.push:
        latest_t = max([item["time"].strftime("%H:%M") for item in fetched])
        git_commit_and_push(latest_t)

if __name__ == "__main__":
    main()
